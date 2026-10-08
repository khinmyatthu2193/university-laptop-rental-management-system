from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Laptop, LaptopAssignment, ManagementStaff, Person, Student, SystemPreference


class PhaseOneSecurityTests(TestCase):
    password = 'A-strong-local-password-2026'

    def create_management_staff(self, role='SystemAdmin', username='operator', permissions=None):
        user = User.objects.create_user(
            username=username,
            email=f'{username}@miit.edu.mm',
            password=self.password,
        )
        profile = ManagementStaff.objects.create(
            auth_user=user,
            name=username.title(),
            username=username,
            outlook_mail=f'{username}@miit.edu.mm',
            position='Admin Staff',
            department='ITSM',
            role=role,
            permissions=permissions or [],
        )
        return user, profile

    def test_root_redirects_to_login(self):
        response = self.client.get('/')
        self.assertRedirects(response, reverse('rental_system:login'))

    def test_operational_pages_require_login(self):
        protected_names = [
            'home', 'student_list', 'staff_list', 'inventory_list',
            'assigned_laptop_list', 'return_laptop_list', 'issue_list',
            'audit_logs',
        ]
        for name in protected_names:
            response = self.client.get(reverse(f'rental_system:{name}'))
            self.assertRedirects(response, reverse('rental_system:login'))

    def test_login_uses_django_authentication_and_email(self):
        user, profile = self.create_management_staff()
        response = self.client.post(reverse('rental_system:login'), {
            'email': profile.outlook_mail,
            'password': self.password,
        })
        self.assertRedirects(response, reverse('rental_system:home'))
        session = self.client.session
        self.assertEqual(int(session['_auth_user_id']), user.id)
        self.assertEqual(session['staff_id'], profile.id)

    def test_public_registration_closes_after_initial_account(self):
        self.create_management_staff()
        response = self.client.get(reverse('rental_system:signin'))
        self.assertRedirects(response, reverse('rental_system:login'))

    def test_login_is_rate_limited_after_five_failures(self):
        _, profile = self.create_management_staff(username='limited')
        for _ in range(5):
            self.client.post(reverse('rental_system:login'), {
                'email': profile.outlook_mail,
                'password': 'wrong-password',
            })
        response = self.client.post(reverse('rental_system:login'), {
            'email': profile.outlook_mail,
            'password': self.password,
        })
        self.assertContains(response, 'Too many failed login attempts', status_code=200)

    def test_inactive_account_cannot_login(self):
        user, profile = self.create_management_staff(username='inactive')
        profile.status = 'Resigned'
        profile.save(update_fields=['status'])
        response = self.client.post(reverse('rental_system:login'), {
            'email': profile.outlook_mail,
            'password': self.password,
        })
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertContains(response, 'Invalid email or password', status_code=200)

    def test_read_only_account_can_view_but_cannot_write(self):
        user, _ = self.create_management_staff(
            role='ReadOnly', username='reporter', permissions=['inventory.view']
        )
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('rental_system:inventory_list')).status_code, 200)
        self.assertEqual(self.client.post(reverse('rental_system:laptop_create'), {}).status_code, 403)

    def test_named_owner_always_has_full_access(self):
        user, profile = self.create_management_staff(
            role='Admin', username='daw_moe_thida', permissions=[]
        )
        self.client.force_login(user)
        self.assertTrue(profile.is_account_owner)
        self.assertEqual(self.client.get(reverse('rental_system:inventory_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('rental_system:assign_new_admin')).status_code, 200)

    def test_permissions_limit_other_accounts_by_module_and_action(self):
        self.create_management_staff(username='daw_moe_thida')
        user, _ = self.create_management_staff(
            role='Admin', username='inventory_viewer', permissions=['inventory.view']
        )
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('rental_system:inventory_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('rental_system:student_list')).status_code, 403)
        self.assertEqual(self.client.post(reverse('rental_system:laptop_create'), {}).status_code, 403)
        self.assertEqual(self.client.get(reverse('rental_system:assign_new_admin')).status_code, 403)

    def test_owner_can_create_account_with_selected_permissions(self):
        user, _ = self.create_management_staff(username='daw_moe_thida')
        self.client.force_login(user)
        response = self.client.post(reverse('rental_system:assign_new_admin'), {
            'name': 'New ITSM Staff',
            'username': 'new_itsm_staff',
            'email': 'new_itsm_staff@miit.edu.mm',
            'position': 'IT Support Officer',
            'department': 'ITSM',
            'password': self.password,
            'confirm_password': self.password,
            'access_inventory': 'manage',
            'access_students': 'view',
        })
        self.assertRedirects(response, reverse('rental_system:assign_new_admin'))
        created = ManagementStaff.objects.get(username='new_itsm_staff')
        self.assertEqual(created.department, 'ITSM')
        self.assertTrue(created.must_change_password)
        self.assertEqual(
            created.permissions,
            ['students.view', 'inventory.view', 'inventory.manage'],
        )

    def test_temporary_password_must_be_changed_before_system_access(self):
        self.create_management_staff(username='daw_moe_thida')
        _user, profile = self.create_management_staff(
            role='Admin', username='temporary_user', permissions=['inventory.view']
        )
        profile.must_change_password = True
        profile.save(update_fields=['must_change_password'])

        response = self.client.post(reverse('rental_system:login'), {
            'email': profile.outlook_mail,
            'password': self.password,
        })
        self.assertRedirects(response, reverse('rental_system:account_settings'))
        self.assertRedirects(
            self.client.get(reverse('rental_system:inventory_list')),
            reverse('rental_system:account_settings'),
        )

        new_password = 'A-different-strong-password-2026'
        response = self.client.post(reverse('rental_system:change_password'), {
            'current_password': self.password,
            'new_password': new_password,
            'confirm_password': new_password,
        })
        self.assertRedirects(response, reverse('rental_system:account_settings'))
        profile.refresh_from_db()
        self.assertFalse(profile.must_change_password)
        self.assertEqual(self.client.get(reverse('rental_system:inventory_list')).status_code, 200)

    def test_only_system_administrator_can_create_accounts(self):
        user, _ = self.create_management_staff(role='Admin', username='admin')
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('rental_system:assign_new_admin')).status_code, 403)

    def test_logout_and_deactivation_reject_get(self):
        user, _ = self.create_management_staff()
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('rental_system:logout')).status_code, 405)
        self.assertEqual(self.client.get(reverse('rental_system:deactivate_account')).status_code, 405)

    def test_preferences_are_persisted(self):
        user, profile = self.create_management_staff()
        self.client.force_login(user)
        response = self.client.post(reverse('rental_system:save_preferences'), {
            'system_name': 'UniKit Production',
            'timezone': 'Asia/Yangon',
            'default_rental_days': '30',
        })
        self.assertRedirects(response, reverse('rental_system:system_preferences'))
        preference = SystemPreference.objects.get(pk=1)
        self.assertEqual(preference.updated_by, profile)
        self.assertEqual(preference.values['system_name'], 'UniKit Production')


class AssignmentIntegrityTests(TestCase):
    def setUp(self):
        self.person = Person.objects.create(
            name='Test Student',
            phone_number='09123456789',
            outlook_mail='student@miit.edu.mm',
            person_type='Student',
        )
        self.laptop = Laptop.objects.create(
            SerialNumber='TEST-001',
            name='Test Model',
            brand='Test Brand',
            processor_gen='Test CPU',
            ram=8,
            storage='256GB SSD',
            for_whom='Student',
        )

    def assignment(self, **overrides):
        today = date.today()
        values = {
            'person': self.person,
            'laptop': self.laptop,
            'issue_date': today,
            'expected_return_date': today + timedelta(days=14),
            'academic_year': str(today.year),
            'assignment_status': 'Issued',
        }
        values.update(overrides)
        return LaptopAssignment.objects.create(**values)

    def test_laptop_cannot_have_two_active_assignments(self):
        self.assignment()
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.assignment(assignment_status='Overdue')

    def test_person_cannot_have_two_active_assignments(self):
        self.assignment()
        second_laptop = Laptop.objects.create(
            SerialNumber='TEST-002',
            name='Second Model',
            brand='Test Brand',
            processor_gen='Test CPU',
            ram=8,
            storage='256GB SSD',
            for_whom='Student',
        )
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.assignment(laptop=second_laptop)

    def test_due_date_cannot_precede_issue_date(self):
        today = date.today()
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.assignment(expected_return_date=today - timedelta(days=1))


class LegacyAssignmentImportTests(TestCase):
    password = 'A-strong-local-password-2026'

    def setUp(self):
        user = User.objects.create_user(
            username='legacy_importer',
            email='legacy_importer@miit.edu.mm',
            password=self.password,
        )
        ManagementStaff.objects.create(
            auth_user=user,
            name='Legacy Importer',
            username='legacy_importer',
            outlook_mail='legacy_importer@miit.edu.mm',
            position='IT Support',
            department='ITSM',
            role='Admin',
            permissions=['assignments.view', 'assignments.manage'],
        )
        self.client.force_login(user)
        self.student = Student.objects.create(
            student_id='2026-MIIT-CSE-001',
            full_name='Test Student',
            email='2026-miit-cse-001@miit.edu.mm',
            phone='',
            major='CSE',
            batch_year=2026,
        )
        self.laptop = Laptop.objects.create(
            SerialNumber='DEVICE-001',
            name='Test Model',
            brand='Test Brand',
            processor_gen='Test CPU',
            ram=8,
            storage='256GB SSD',
            for_whom='Student',
        )

    def workbook_upload(self, rows):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            'No.', 'Name', 'Roll Number', 'Device Number', 'Issue Date',
            'Return Date', 'Sign', 'Phone No', 'Remark',
        ])
        for row in rows:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            'legacy_assignments.xlsx',
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_downloaded_template_uses_english_headers(self):
        response = self.client.get(reverse('rental_system:legacy_assignments_template'))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        self.assertEqual(headers, [
            'No.', 'Name', 'Roll Number', 'Device Number', 'Issue Date',
            'Return Date', 'Sign', 'Phone No', 'Remark',
        ])

    def test_imports_returned_legacy_assignment_and_metadata(self):
        issue_date = date(2025, 1, 10)
        return_date = date(2025, 2, 10)
        upload = self.workbook_upload([[
            1, self.student.full_name, self.student.student_id,
            self.laptop.SerialNumber, issue_date, return_date,
            'Signed', '09123456789', 'Returned in good condition',
        ]])
        response = self.client.post(
            reverse('rental_system:import_legacy_assignments'),
            {'file': upload},
        )
        self.assertRedirects(response, reverse('rental_system:assigned_laptop_list'))
        assignment = LaptopAssignment.objects.get()
        self.assertEqual(assignment.assignment_status, 'Returned')
        self.assertEqual(assignment.actual_return_date, return_date)
        self.assertEqual(assignment.legacy_sign, 'Signed')
        self.assertEqual(assignment.remark, 'Returned in good condition')
        self.student.refresh_from_db()
        self.assertIsNone(self.student.laptop_id)
        self.assertEqual(self.student.phone, '09123456789')
        self.assertEqual(self.student.rental_status, 'Returned')

    def test_imports_blank_return_date_as_current_assignment(self):
        today = date.today()
        upload = self.workbook_upload([[
            1, self.student.full_name, self.student.student_id,
            self.laptop.SerialNumber, today, None, '', '', '',
        ]])
        response = self.client.post(
            reverse('rental_system:import_legacy_assignments'),
            {'file': upload},
        )
        self.assertRedirects(response, reverse('rental_system:assigned_laptop_list'))
        assignment = LaptopAssignment.objects.get()
        self.assertEqual(assignment.assignment_status, 'Issued')
        self.assertIsNone(assignment.actual_return_date)
        self.student.refresh_from_db()
        self.laptop.refresh_from_db()
        self.assertEqual(self.student.laptop_id, self.laptop.id)
        self.assertEqual(self.laptop.status, 'Assigned')
        self.assertEqual(self.student.rental_status, 'Assigned')

    def test_imports_student_without_device_as_not_requested(self):
        upload = self.workbook_upload([[
            1, 'New Student', '2025-MIIT-CSE-099',
            None, None, None, '', '09111111111', '',
        ]])
        response = self.client.post(
            reverse('rental_system:import_legacy_assignments'),
            {'file': upload},
        )
        self.assertRedirects(response, reverse('rental_system:assigned_laptop_list'))
        student = Student.objects.get(student_id='2025-MIIT-CSE-099')
        self.assertEqual(student.full_name, 'New Student')
        self.assertEqual(student.major, 'CSE')
        self.assertEqual(student.batch_year, 2025)
        self.assertEqual(student.email, '2025-miit-cse-099@miit.edu.mm')
        self.assertEqual(student.rental_status, 'Not Requested')
        self.assertIsNone(student.laptop_id)
        self.assertFalse(LaptopAssignment.objects.exists())

    def test_invalid_workbook_imports_nothing(self):
        upload = self.workbook_upload([[
            1, 'Unknown Student', 'UNKNOWN-ROLL', self.laptop.SerialNumber,
            date.today(), None, '', '', '',
        ]])
        response = self.client.post(
            reverse('rental_system:import_legacy_assignments'),
            {'file': upload},
        )
        self.assertRedirects(response, reverse('rental_system:assigned_laptop_list'))
        self.assertFalse(LaptopAssignment.objects.exists())

    def test_finds_legacy_table_on_later_sheet_and_parses_excel_date(self):
        workbook = Workbook()
        unrelated = workbook.active
        unrelated.title = 'Other Data'
        unrelated.append(['Roll No.', "Student's Name"])
        unrelated.append(['2025_0001', 'Different List'])
        worksheet = workbook.create_sheet('ECE Roll Number')
        worksheet.append(['Legacy title'])
        worksheet.append(['Another title'])
        worksheet.append([])
        worksheet.append([
            'စဉ်', 'Extra Roll', 'အမည်', 'ပထမနှစ်သင်တန်း\nခုံအမှတ်',
            'စက်နံပါတ်', 'Issue Date', 'Return Date', 'Sign', 'Phone No', 'Remark',
        ])
        worksheet.append([
            1, '25_0001', self.student.full_name, self.student.student_id,
            self.laptop.SerialNumber, 45669, None, '', '09123456789', '',
        ])
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            'multi_sheet_legacy.xlsx',
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        response = self.client.post(
            reverse('rental_system:import_legacy_assignments'),
            {'file': upload},
        )

        self.assertRedirects(response, reverse('rental_system:assigned_laptop_list'))
        assignment = LaptopAssignment.objects.get()
        self.assertEqual(assignment.issue_date, date(2025, 1, 12))
        self.assertEqual(assignment.assignment_status, 'Overdue')
